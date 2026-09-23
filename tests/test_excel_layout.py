import pytest
from xml.etree import ElementTree
from zipfile import ZipFile
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from src.config import Export
from src.export.excel_exporter import EN_WIDTHS_PX, LEMMA_COLUMNS, english_columns, write_excel


@pytest.mark.parametrize("description", [" credit ", "\tcredit\n", "credit", "   "])
def test_description_is_valid_opc_metadata(tmp_path, description):
    target = tmp_path / "metadata.xlsx"
    write_excel(target, [("Test", ["Text"], [[" content "]])], Export(), description=description)
    with ZipFile(target) as archive:
        core = ElementTree.fromstring(archive.read("docProps/core.xml"))
    assert all("{http://www.w3.org/XML/1998/namespace}space" not in node.attrib for node in core.iter())
    book = load_workbook(target)
    assert (book.properties.description or "") == description.strip()
    assert book.active["A2"].value == " content "
    book.close()


def test_empty_optional_columns_only_and_no_input_mutation():
    headers = ["Лемма", "Род", "Регулярность", "Перевод на английский", "Перевод на русский"]
    rows = [["watch", None, "", " ", ""], ["book", float("nan"), "regular", "", ""]]
    filtered, values = english_columns(headers, rows)
    assert filtered == ["Лемма", "Регулярность", "Перевод на русский"]
    assert values == [["watch", "", ""], ["book", "regular", ""]]
    assert len(headers) == len(rows[0]) == 5
    assert english_columns(headers, [])[0] == ["Лемма", "Перевод на русский"]


@pytest.mark.parametrize("language", ["en", "es"])
def test_workbook_layout_language_scope_and_widths(tmp_path, language):
    headers = LEMMA_COLUMNS
    values = {"Лемма": "watch", "Количество контекстных блоков": 3, "Перевод на русский": "часы"}
    rows = [[values.get(header, "") for header in headers]]
    target = tmp_path / "dictionary.xlsx"
    write_excel(target, [("Леммы", headers, rows)], Export(), language)
    book = load_workbook(target)
    sheet = book.active
    exported = [cell.value for cell in sheet[1]]
    index = exported.index("Количество контекстных блоков")
    assert exported[index - 1 : index + 2] == [
        "Специфичность",
        "Количество контекстных блоков",
        "Транскрипция",
    ]
    assert sheet.cell(2, index + 1).value == 3
    for header in ("Род", "Группа спряжения", "Регулярность", "Перевод на английский"):
        assert (header in exported) == (language == "es")
    assert sheet.cell(2, exported.index("Перевод на русский") + 1).value == "часы"
    if language == "en":
        for header, pixels in EN_WIDTHS_PX.items():
            column = get_column_letter(exported.index(header) + 1)
            stored_width = sheet.column_dimensions[column].width
            # ISO 29500 conversion from the saved OOXML width to pixels.
            assert int(((256 * stored_width + int(128 / 7)) / 256) * 7) == pixels
    assert sheet.auto_filter.ref == sheet.dimensions
    assert sheet.freeze_panes == "A2"
    book.close()

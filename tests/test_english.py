import json

from src.cards.docx_renderer import render
from src.cards.formatter import english_back_forms
from src.cards.known_words import read_known_words
from src.cards.models import CardEntry
from src.cards.selector import select_cards
from src.grammar.english import enrich, plural
from src.languages import get_profile
from src.postprocessing.processor import process
from src.translation.openai_translator import parse_response


def test_english_profile_and_plural_and_verb_forms():
    assert get_profile("en").wordfreq_language == "en"
    assert plural("city") == "cities"
    assert plural("child") == "children"

    class Config:
        enabled = True

    grammar, _ = enrich(
        [
            {"id": "go", "lemma": "go", "pos": "VERB"},
            {"id": "book", "lemma": "book", "pos": "NOUN"},
        ],
        [],
        Config(),
    )
    assert grammar[0].grammar_forms == "go / goes / went / gone / going"
    assert grammar[1].grammar_forms == "book / books"


def test_english_translation_contract_has_only_russian():
    assert parse_response(json.dumps({"entries": [{"id": "a", "ru": "идти"}]}), {"a"}, "en") == [
        {"id": "a", "ru": "идти"}
    ]


def test_english_card_back_omits_english_translation(tmp_path):
    target = tmp_path / "cards.docx"
    template = __import__("pathlib").Path(__file__).parents[1] / "table.docx"
    render(
        [
            CardEntry(
                "go", "VERB", 1, "go / goes / went / gone / going", "go", "ɡoʊ", "идти", "to go", {"gone": 2}
            )
        ],
        template,
        target,
        type(
            "Config",
            (),
            {
                "cards_per_page": 24,
                "rows": 8,
                "columns": 3,
                "front_font_size": 14,
                "back_font_size": 12,
                "max_forms": 10,
                "long_translation_threshold": 35,
                "long_translation_font_size": 10,
                "long_forms_threshold": 90,
                "long_forms_font_size": 10,
            },
        )(),
        "en",
    )
    from docx import Document

    document = Document(target)
    assert document.tables[0].cell(0, 0).text == "go"
    assert [p.text for p in document.tables[1].cell(7, 0).paragraphs] == [
        "/ɡoʊ/",
        "идти",
        "goes, went, gone, going",
    ]


def test_english_card_front_is_learning_form_and_back_combines_forms():
    entry = CardEntry(
        "be",
        "AUX",
        1,
        "be / am / are / is / was / were / been / being",
        "be",
        "biː",
        "быть",
        "to be",
        {"is": 8, "was": 7, "been": 5, "being": 3, "were": 2, "are": 1, "am": 1},
    )
    assert english_back_forms(entry) == ["am", "are", "is", "was", "were", "been", "being"]


def test_english_postprocessing_removes_detached_apostrophe_suffixes():
    lemmas = [
        {"id": "walk", "lemma": "walk", "pos": "VERB", "count": 3},
        {"id": "ed", "lemma": "’ed", "pos": "AUX", "count": 2},
        {"id": "ll", "lemma": "'ll", "pos": "AUX", "count": 1},
    ]
    forms = [
        {"id": "walk-form", "form": "walk", "lemma": "walk", "pos": "VERB", "count": 3},
        {"id": "ed-form", "form": "’ed", "lemma": "’ed", "pos": "AUX", "count": 2},
        {"id": "ll-form", "form": "'ll", "lemma": "'ll", "pos": "AUX", "count": 1},
    ]
    clean_lemmas, clean_forms = process(lemmas, forms, "en")
    assert [row["lemma"] for row in clean_lemmas] == ["walk"]
    assert [row["form"] for row in clean_forms] == ["walk"]
    assert clean_lemmas[0]["cumulative_coverage"] == 1.0


def test_english_postprocessing_normalizes_terminal_period(monkeypatch):
    monkeypatch.setattr("src.postprocessing.processor.get_reference_frequency", lambda *_: 0.01)
    lemmas = [
        {"id": "mrs", "lemma": "mrs.", "pos": "PROPN", "count": 2, "observed_forms": {"mrs.": 2}},
    ]
    forms = [
        {
            "id": "mrs-form",
            "form": "mrs.",
            "lemma": "mrs.",
            "pos": "PROPN",
            "count": 2,
            "reference_frequency": 0.0,
        },
    ]
    clean_lemmas, clean_forms = process(lemmas, forms, "en")
    assert clean_lemmas[0]["lemma"] == "mrs"
    assert clean_lemmas[0]["observed_forms"] == {"mrs": 2}
    assert clean_forms[0]["form"] == clean_forms[0]["lemma"] == "mrs"


def test_english_postprocessing_removes_negative_contraction_stems():
    lemmas = [
        {"id": "would", "lemma": "would", "pos": "AUX", "count": 2},
        {"id": "wouldn", "lemma": "wouldn", "pos": "AUX", "count": 2},
    ]
    forms = [
        {"id": "would", "form": "would", "lemma": "would", "pos": "AUX", "count": 2},
        {"id": "wouldn", "form": "wouldn", "lemma": "wouldn", "pos": "AUX", "count": 2},
    ]
    clean_lemmas, clean_forms = process(lemmas, forms, "en")
    assert [row["lemma"] for row in clean_lemmas] == ["would"]
    assert [row["form"] for row in clean_forms] == ["would"]


def test_known_dictionary_excludes_matching_lemmas(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "my_dictionary_en.xlsx"
    workbook = Workbook()
    workbook.active.append(["Known words"])
    workbook.active.append(["Go"])
    workbook.save(path)
    workbook.close()
    cards, _ = select_cards(
        {
            "lemmas": [
                {
                    "id": "go",
                    "lemma": "go",
                    "pos": "VERB",
                    "rank": 1,
                    "count": 3,
                    "cumulative_coverage": 1.0,
                    "observed_forms": {},
                },
            ],
            "forms": [],
            "lemma_grammar": [],
            "ipa": [],
            "translations": [],
        },
        100,
        language="en",
        known_words=read_known_words(path),
    )
    assert cards == []

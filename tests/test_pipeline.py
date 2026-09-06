import json
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

from src.manifest import DEPENDENCIES
from src.pipeline import Pipeline
from src.storage import read_rows, write_rows


@pytest.fixture
def project(tmp_path, monkeypatch, observations):
    config_path = tmp_path / "config.yaml"
    config = {
        "nlp": {"model": "test"},
        "pronunciation": {"enabled": False},
        "translation": {"enabled": False},
        "cards": {"enabled": False},
    }
    config_path.write_text(yaml.safe_dump(config), encoding="utf-8")
    source = tmp_path / "book.txt"
    source.write_text("Juan vino a casa.\nEl vino era bueno.\n", encoding="utf-8")

    def fake_nlp(source, target, config):
        write_rows(target, observations)

    monkeypatch.setattr("src.nlp.processor.run", fake_nlp)
    monkeypatch.setattr(
        "src.reference_frequency.calculator.get_reference_frequency", lambda form, language, lookup=None: 0.01
    )
    pipeline = Pipeline(config_path, source)
    pipeline.run("run-all")
    return config_path, source


def generations(pipeline):
    return {k: v["generation"] for k, v in pipeline.manifest.data["stages"].items()}


def test_no_repeated_work(project):
    path, source = project
    pipeline = Pipeline(path, source)
    before = generations(pipeline)
    pipeline.run("run-all")
    assert before == generations(pipeline)
    assert all(state == "OK" for state, _ in pipeline.status().values())


@pytest.mark.parametrize(
    "section,key,value,stale",
    [
        ("pronunciation", "language", "es-419", {"ipa", "validate", "export", "cards"}),
        ("translation", "prompt_version", "2.0", {"translate", "validate", "export", "cards"}),
        (
            "translation",
            "cumulative_coverage_limit",
            85,
            {"translate", "validate", "export", "cards"},
        ),
        ("grammar", "enabled", False, {"grammar", "validate", "export", "cards"}),
        ("export", "font_size", 12, {"export", "cards"}),
        ("nlp", "model", "another", set(DEPENDENCIES) - {"preprocess"}),
        ("filters", "include_proper_nouns", False, set(DEPENDENCIES) - {"preprocess", "nlp"}),
    ],
)
def test_config_invalidation(project, section, key, value, stale):
    path, _ = project
    cfg = yaml.safe_load(path.read_text(encoding="utf-8"))
    cfg.setdefault(section, {})[key] = value
    path.write_text(yaml.safe_dump(cfg), encoding="utf-8")
    statuses = Pipeline(path).status()
    assert {s for s, (state, _) in statuses.items() if state == "OUTDATED"} == stale
    first = (
        "translate"
        if section == "translation"
        else "ipa"
        if section == "pronunciation"
        else "aggregate"
        if section == "filters"
        else section
    )
    assert key in statuses[first][1]


def test_source_change_and_artifact_tampering(project):
    path, source = project
    source.write_text("Otra casa.", encoding="utf-8")
    statuses = Pipeline(path).status()
    assert all(statuses[s][0] == "OUTDATED" for s in DEPENDENCIES)
    assert statuses["source"][0] == "CHANGED"


def test_force_ipa_only_then_export(project):
    path, _ = project
    pipeline = Pipeline(path)
    before = generations(pipeline)
    pipeline.run("ipa", force=True)
    after = generations(pipeline)
    assert {s for s in after if before[s] != after[s]} == {"ipa"}
    assert pipeline.status()["postprocess"][0] == "OK"
    pipeline.run("export")
    after = generations(pipeline)
    assert {s for s in after if before[s] != after[s]} == {"ipa", "validate", "export"}


def test_missing_artifact_refuses_expensive_automatic_run(project):
    path, _ = project
    pipeline = Pipeline(path)
    pipeline.outputs["nlp"][0].unlink()
    assert pipeline.status()["aggregate"][0] == "OUTDATED"
    with pytest.raises(RuntimeError, match="зависимость"):
        pipeline.run("export")


def test_corruption_and_failed_stage(project, monkeypatch):
    path, _ = project
    pipeline = Pipeline(path)
    pipeline.outputs["ipa"][0].write_bytes(b"changed")
    assert pipeline.status()["ipa"][0] == "OUTDATED"

    def fail(stage):
        raise RuntimeError("failure")

    monkeypatch.setattr(pipeline, "execute", fail)
    with pytest.raises(RuntimeError, match="failure"):
        pipeline.run("ipa")
    assert Pipeline(path).status()["ipa"][0] == "FAILED"


def test_empty_book_pipeline(project, monkeypatch):
    path, source = project
    source.write_text("", encoding="utf-8")
    monkeypatch.setattr("src.nlp.processor.run", lambda source, target, cfg: write_rows(target, []))
    pipeline = Pipeline(path)
    pipeline.run("run-all")
    assert read_rows(pipeline.outputs["aggregate"][0]) == []
    book = load_workbook(pipeline.output / "spanish_frequency_dictionary.xlsx")
    assert book["Леммы"].max_row == 1
    book.close()


def test_excel_csv_and_manifest(project):
    path, _ = project
    pipeline = Pipeline(path)
    book = load_workbook(pipeline.output / "spanish_frequency_dictionary.xlsx")
    assert book.sheetnames == ["Леммы", "Словоформы", "Проверка"]
    for sheet in book:
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref
        assert sheet["A1"].font.bold
        assert sheet["A1"].alignment.wrap_text
    sheet = book["Леммы"]
    assert sheet["D2"].data_type == "n"
    assert sheet["E2"].number_format == "0.00%"
    assert sheet["F2"].value > 0
    assert sheet["R2"].value is None  # Disabled translation stays blank.
    book.close()
    assert (pipeline.output / "lemmas.csv").read_text(encoding="utf-8").startswith("Ранг,Лемма")
    manifest = json.loads((pipeline.work / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["source"]["sha256"]
    assert all(
        r["completed"] and r["config_hash"] and (r["outputs"] or stage == "cards")
        for stage, r in manifest["stages"].items()
    )


def test_excel_formula_injection(tmp_path):
    from src.config import Export
    from src.export.excel_exporter import write_excel

    path = tmp_path / "safe.xlsx"
    write_excel(path, [("Текст", ["Слово"], [["=1+1"], ["+cmd"], ["@test"]])], Export())
    book = load_workbook(path, data_only=False)
    assert book.active["A2"].value == "=1+1"
    assert book.active["A2"].data_type == "s"
    book.close()


def test_prompt_code_hash_invalidation(project, monkeypatch):
    path, _ = project
    import src.pipeline as module

    original = module.file_hash

    def changed_hash(path):
        return "changed" if Path(path).name == "prompts.py" else original(path)

    monkeypatch.setattr(module, "file_hash", changed_hash)
    statuses = Pipeline(path).status()
    assert {s for s, (state, _) in statuses.items() if state == "OUTDATED"} == {
        "translate",
        "validate",
        "export",
        "cards",
    }

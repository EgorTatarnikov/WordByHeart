import pytest

from src.cards.formatter import english_back_forms
from src.cards.selector import select_cards
from src.export.excel_exporter import make_tables
from src.lemma_groups import group_lemmas
from src.translation.selection import select_lemmas_by_cumulative_coverage
from src.config import Export
from src.export.excel_exporter import write_excel
from openpyxl import load_workbook


def sample():
    lemmas = [
        {
            "id": "verb",
            "lemma": "watch",
            "pos": "VERB",
            "count": 6,
            "observed_forms": {"watch": 4, "watches": 2},
        },
        {"id": "other", "lemma": "book", "pos": "NOUN", "count": 10, "observed_forms": {"book": 10}},
        {
            "id": "noun",
            "lemma": "watch",
            "pos": "NOUN",
            "count": 5,
            "observed_forms": {"watch": 3, "watches": 2},
        },
    ]
    forms = [
        {
            "id": r["id"] + f,
            "lemma": r["lemma"],
            "pos": r["pos"],
            "form": f,
            "count": n,
            "reference_frequency": 0.01,
        }
        for r in lemmas
        for f, n in r["observed_forms"].items()
    ]
    for rank, row in enumerate(lemmas + forms, 1):
        row.update(rank=rank, share=row["count"] / 21, cumulative_coverage=1, contexts=[], chunk_count=1)
    grammar = [
        {
            "id": r["id"],
            "learning_form": r["lemma"],
            "grammar_forms": "",
            "grammatical_gender": "",
            "infinitive": "",
            "conjugation_group": "",
            "regularity": "",
            "morph_description": "",
        }
        for r in lemmas + forms
    ]
    grammar[0]["grammar_forms"] = "watch / watches / watched / watching"
    grammar[2]["grammar_forms"] = "watch / watches"
    translations = [
        {
            "id": r["id"],
            "ru": "смотреть, наблюдать" if r["id"] == "verb" else "часы" if r["id"] == "noun" else "книга",
            "en": "",
        }
        for r in lemmas + forms
    ]
    return {
        "lemmas": lemmas,
        "forms": forms,
        "lemma_grammar": grammar[:3],
        "form_grammar": grammar[3:],
        "ipa": [],
        "translations": translations,
        "validation": [],
    }


def test_group_frequency_selection_and_unique_reference():
    data = sample()
    groups = group_lemmas(data["lemmas"], data["forms"])
    assert groups[0]["lemma"] == "watch"
    assert groups[0]["count"] == 11
    assert groups[0]["reference_frequency"] == pytest.approx(0.02)
    selection = select_lemmas_by_cumulative_coverage(data["lemmas"], data["forms"], 50, 1e9)
    assert selection.eligible_lemma_ids == {"verb", "noun"}
    assert selection.actual_coverage == pytest.approx(11 / 21)
    assert len(selection.eligible_form_ids) == 4


@pytest.mark.parametrize("language", ["en", "es"])
def test_export_preserves_frequency_and_specificity(tmp_path, language):
    target = tmp_path / f"{language}_dictionary.xlsx"
    write_excel(target, make_tables(sample()), Export(), language)
    book = load_workbook(target)
    sheet = book["Леммы"]
    headers = [cell.value for cell in sheet[1]]
    assert "Доля от всех слов" not in headers
    assert headers.count("Относительная частота в книге") == 1
    for row, count, reference in ((2, 11, 0.02), (3, 11, 0.02), (4, 10, 0.01)):
        frequency = sheet.cell(row, headers.index("Относительная частота в книге") + 1)
        specificity = sheet.cell(row, headers.index("Специфичность") + 1)
        assert frequency.value == pytest.approx(count / 21)
        assert frequency.number_format == "0.00%"
        wordfreq = sheet.cell(row, headers.index("Частота Wordfreq") + 1)
        assert wordfreq.value == pytest.approx(reference)
        assert wordfreq.number_format == "0.000000%"
        assert specificity.value == pytest.approx((count / 21) / reference)
        assert specificity.number_format == "0.000000"
    forms = book["Словоформы"]
    wordfreq_column = [cell.value for cell in forms[1]].index("Частота Wordfreq") + 1
    for row in range(2, forms.max_row + 1):
        assert forms.cell(row, wordfreq_column).value == pytest.approx(0.01)
        assert forms.cell(row, wordfreq_column).number_format == "0.000000%"
    book.close()


def test_specificity_adds_entire_group_and_counts_actual_coverage():
    data = sample()
    # book leads the frequency prefix; watch qualifies only by combined count.
    data["lemmas"][1]["count"] = 30
    selection = select_lemmas_by_cumulative_coverage(data["lemmas"], data["forms"], 50, 10, 11)
    assert selection.eligible_lemma_ids == {"verb", "noun", "other"}
    assert selection.actual_coverage == 1


def test_single_compact_card_and_adjacent_export_rows():
    data = sample()
    cards, _ = select_cards(data, 50, 1e9, language="en")
    assert len(cards) == 1
    assert cards[0].translation_ru == "смотреть, наблюдать, часы"
    assert cards[0].observed_forms == {"watch": 7, "watches": 4}
    assert english_back_forms(cards[0]) == ["watches", "watched", "watching"]
    _, headers, rows = make_tables(data)[0]
    assert [r[1] for r in rows] == ["watch", "watch", "book"]
    assert [r[headers.index("Всего вхождений леммы")] for r in rows] == [11, 11, 10]
    assert [r[headers.index("Ранг")] for r in rows] == [1, 1, 2]
    assert headers[3:9] == [
        "Суммарное число вхождений",
        "Всего вхождений леммы",
        "Кумулятивное покрытие",
        "Относительная частота в книге",
        "Частота Wordfreq",
        "Специфичность",
    ]
    assert len(headers) == 21
    assert "Доля от всех слов" not in headers
    assert all(len(row) == len(headers) for row in rows)
    for row in rows[:2]:
        assert row[4:9] == pytest.approx([11, 11 / 21, 11 / 21, 0.02, (11 / 21) / 0.02])
    assert rows[-1][headers.index("Кумулятивное покрытие")] == 1
    cards, _ = select_cards(data, 100, language="en", known_words={"watch"})
    assert [c.lemma for c in cards] == ["book"]

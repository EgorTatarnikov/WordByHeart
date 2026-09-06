from dataclasses import replace
from pathlib import Path

import pytest
from docx import Document

from src.cards.docx_renderer import back_position, render, table_order
from src.cards.formatter import (
    extract_forms_from_front,
    filter_observed_forms_for_back,
    format_observed_forms,
    get_card_front_text,
    get_forms_font_size,
    get_translation_font_size,
    select_observed_forms,
)
from src.cards.models import CardEntry
from src.cards.selector import select_cards
from src.config import Cards


def card(rank=1, forms=None):
    return CardEntry(
        lemma="casa",
        pos="NOUN",
        rank=rank,
        grammatical_forms="la casa / las casas",
        learning_form="la casa",
        ipa="kasa",
        translation_ru="дом",
        translation_en="house",
        observed_forms=forms or {"casa": 4, "casas": 2},
    )


def test_front_text_priority_and_empty_values():
    assert get_card_front_text(card()) == "la casa / las casas"
    assert get_card_front_text(replace(card(), grammatical_forms=None)) == "la casa"
    assert get_card_front_text(replace(card(), grammatical_forms="  ", learning_form="")) == "casa"


def test_observed_forms_limit_counts_and_threshold():
    forms = {f"forma{i}": 20 - i for i in range(15)}
    selected = select_observed_forms(forms)
    assert selected == [f"forma{i}" for i in range(10)]
    assert format_observed_forms({"x": 1} if False else selected).count("(") == 0
    assert get_forms_font_size("x" * 90, 12, 90, 10) == 12
    assert get_forms_font_size("x" * 91, 12, 90, 10) == 10
    assert get_translation_font_size("x" * 35, 12, 35, 10) == 12
    assert get_translation_font_size("x" * 36, 12, 35, 10) == 10


@pytest.mark.parametrize(
    ("front", "observed", "expected"),
    [
        ("Harry", {"Harry": 1}, []),
        ("la profesora / las profesoras", {"profesora": 2, "profesoras": 1}, []),
        ("la casa / las casas", {"casa": 2, "casas": 1}, []),
        (
            "pequeño / pequeña / pequeños / pequeñas",
            {"pequeño": 4, "pequeña": 3, "pequeños": 2, "pequeñas": 1},
            [],
        ),
        (
            "tener",
            {"tiene": 40, "tenía": 22, "tengo": 17, "tener": 15, "tuvo": 11},
            ["tiene", "tenía", "tengo", "tuvo"],
        ),
        (
            "el profesor / los profesores",
            {"profesor": 24, "profesores": 8, "profesora": 5, "profesoras": 2},
            ["profesora", "profesoras"],
        ),
        ("Harry", {"harry": 1}, []),
    ],
)
def test_observed_forms_duplicated_on_front_are_removed(front, observed, expected):
    assert filter_observed_forms_for_back(observed, front) == expected


def test_front_form_extraction_and_filtering_happen_before_max_forms():
    assert extract_forms_from_front("la profesora / las profesoras") == {"profesora", "profesoras"}
    observed = {f"forma{i}": 20 - i for i in range(15)}
    front = " / ".join(f"forma{i}" for i in range(5))
    assert filter_observed_forms_for_back(observed, front, max_forms=10) == [
        f"forma{i}" for i in range(5, 15)
    ]


def test_no_observed_forms_paragraph_and_font_threshold_after_filtering(tmp_path):
    template = Path(__file__).parents[1] / "table.docx"
    target = tmp_path / "cards.docx"
    no_forms = replace(card(), observed_forms={"casa": 4, "casas": 2})
    render([no_forms], template, target, Cards())
    back = Document(target).tables[1].cell(7, 0)
    assert [paragraph.text for paragraph in back.paragraphs] == ["/kasa/", "дом", "house"]

    long_forms = {"a" * 30: 4, "b" * 30: 3, "c" * 30: 2, "d": 1}
    front = " / ".join(["a" * 30, "b" * 30])
    forms_text = format_observed_forms(filter_observed_forms_for_back(long_forms, front))
    assert len(format_observed_forms(list(long_forms))) > 90
    assert len(forms_text) <= 90
    assert get_forms_font_size(forms_text, 12, 90, 10) == 12


def test_long_translations_use_ten_point_font(tmp_path):
    template = Path(__file__).parents[1] / "table.docx"
    target = tmp_path / "cards.docx"
    long_ru = "р" * 36
    long_en = "e" * 36
    render([replace(card(), translation_ru=long_ru, translation_en=long_en)], template, target, Cards())
    paragraphs = Document(target).tables[1].cell(7, 0).paragraphs
    runs = [next(run for run in paragraph.runs if run.text) for paragraph in paragraphs[:3]]
    assert runs[0].font.size.pt == 12
    assert runs[1].font.size.pt == 10
    assert runs[2].font.size.pt == 10


@pytest.mark.parametrize(
    ("count", "expected"),
    [
        (24, [("front", 0), ("back", 0)]),
        (25, [("front", 0), ("front", 1), ("back", 1), ("back", 0)]),
        (72, [("front", 0), ("front", 1), ("front", 2), ("back", 2), ("back", 1), ("back", 0)]),
    ],
)
def test_table_order_and_back_mapping(count, expected):
    sheets = (count + 23) // 24
    assert table_order(sheets) == expected
    assert back_position(0, 0, 8) == (7, 0)
    assert back_position(0, 2, 8) == (7, 2)
    assert back_position(7, 1, 8) == (0, 1)


def test_render_uses_template_and_mirrors_back_rows(tmp_path):
    template = Path(__file__).parents[1] / "table.docx"
    target = tmp_path / "cards.docx"
    assert render([card(rank=i + 1) for i in range(25)], template, target, Cards()) == 2
    document = Document(target)
    assert len(document.tables) == 4
    assert document.tables[0].cell(0, 0).text == "la casa / las casas"
    assert document.tables[3].cell(7, 0).paragraphs[0].text == "/kasa/"
    assert document.tables[1].cell(0, 0).text == "la casa / las casas"
    assert document.tables[2].cell(7, 0).paragraphs[0].text == "/kasa/"
    assert not document.tables[1].cell(0, 1).text
    # The final one-point paragraph keeps Word from adding a normal-height line
    # after the last table and spilling it onto a blank page.
    final_paragraph = document.paragraphs[-1]
    assert final_paragraph.runs[0].font.size.pt == 1


def test_cards_use_the_shared_coverage_selector_and_use_placeholders_for_missing_data():
    lemmas = [
        {
            "id": "a",
            "lemma": "casa",
            "pos": "NOUN",
            "rank": 1,
            "count": 9,
            "cumulative_coverage": 0.9,
            "observed_forms": {"casa": 9},
        },
        {
            "id": "b",
            "lemma": "venir",
            "pos": "VERB",
            "rank": 2,
            "count": 1,
            "cumulative_coverage": 1.0,
            "observed_forms": {"vino": 1},
        },
    ]
    data = {
        "lemmas": lemmas,
        "forms": [],
        "lemma_grammar": [
            {"id": "a", "learning_form": "la casa", "grammar_forms": ""},
            {"id": "b", "learning_form": "venir", "grammar_forms": ""},
        ],
        "ipa": [{"text": "casa", "ipa": "kasa"}, {"text": "venir", "ipa": "benir"}],
        "translations": [
            {"id": "a", "ru": "дом", "en": "house"},
            {"id": "b", "ru": "приходить", "en": "come"},
        ],
    }
    cards, selection = select_cards(data, 90)
    assert [entry.lemma for entry in cards] == ["casa"]
    assert selection.actual_coverage == 0.9
    data["translations"] = []
    cards, _ = select_cards(data, 90)
    assert cards[0].translation_ru == "—"
    assert cards[0].translation_en == "—"
    data["ipa"] = []
    cards, _ = select_cards(data, 90)
    assert cards[0].ipa == "-"

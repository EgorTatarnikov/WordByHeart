"""Optional user-maintained known-word lists for card selection."""

from pathlib import Path

from openpyxl import load_workbook

from src.cards.formatter import normalize_card_form


def read_known_words(path: Path) -> set[str]:
    if not path.is_file():
        raise ValueError(f"Known dictionary not found: {path}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        return {
            normalize_card_form(row[0])
            for row in sheet.iter_rows(min_col=1, max_col=1, values_only=True)
            if row and isinstance(row[0], str) and row[0].strip()
        }
    finally:
        workbook.close()

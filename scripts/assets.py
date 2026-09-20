"""Validate supplied user files without creating or modifying them."""
from pathlib import Path


def check_assets(root):
    from docx import Document
    from openpyxl import load_workbook

    root = Path(root)
    for name in ("table.docx", "my_dictionary_en.xlsx", "my_dictionary_es.xlsx"):
        if not (root / name).is_file():
            raise FileNotFoundError(f"Restore the original {name} in {root}.")
    doc = Document(root / "table.docx")
    if len(doc.tables) < 2 or any(len(t.rows) != 8 or len(t.columns) != 3 for t in doc.tables[:2]):
        raise ValueError("Card template must contain two 8 x 3 tables.")
    for language in ("en", "es"):
        workbook = load_workbook(root / f"my_dictionary_{language}.xlsx", read_only=True)
        workbook.close()


if __name__ == "__main__":
    check_assets(Path(__file__).resolve().parents[1])
